# -*- coding: utf-8 -*-
"""
Generate test_valid_exam.docx and test_invalid_exam.docx
"""
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor

sys.stdout.reconfigure(encoding='utf-8')

def create_valid_exam():
    doc = Document()
    doc.add_paragraph("[EXAM_BEGIN]")
    doc.add_paragraph("[EXAM_TITLE: TOEIC READING TEST - MA DE VALID TEST]")
    doc.add_paragraph("[EXAM_DURATION: 75]")
    doc.add_paragraph("[TOTAL_QUESTIONS: 5]")
    doc.add_paragraph("[PART: 5]")
    doc.add_paragraph("[SHUFFLE_QUESTIONS: TRUE]")
    doc.add_paragraph("[SHUFFLE_ANSWERS: TRUE]")

    questions = [
        (101, "Ms. Carpenter will be attending the conference with ------- marketing team.",
         "she", "her", "hers", "herself", "B", "Cần tính từ sở hữu 'her' đứng trước cụm danh từ 'marketing team'."),
        (102, "------- the last ten years, Bay City's population has grown by about 27 percent.",
         "As", "Against", "During", "Below", "C", "Cụm 'During the last ten years' chỉ khoảng thời gian xảy ra hành động."),
        (103, "Please congratulate Alan Schmit, ------- of the Leadership Award in Nursing at Knoll Hospital.",
         "won", "wins", "winning", "winner", "D", "Danh từ chỉ người 'winner' làm đồng vị ngữ giải thích cho Alan Schmit."),
        (104, "Ireland's largest software producer will ------- be opening a large facility in Cork.",
         "soon", "such", "ever", "like", "A", "Trạng từ 'soon' đứng giữa trợ động từ will và be opening."),
        (105, "Ashburn Bank's online service has been in high demand -------.",
         "lateness", "later", "lately", "latest", "C", "Trạng từ 'lately' nhận biết thì hiện tại hoàn thành.")
    ]

    for num, text, a, b, c, d, corr, hint in questions:
        doc.add_paragraph(f"[QUESTION: {num}]")
        doc.add_paragraph(f"{num}. {text}")
        doc.add_paragraph(f"[A] {a}")
        doc.add_paragraph(f"[B] {b}")
        doc.add_paragraph(f"[C] {c}")
        doc.add_paragraph(f"[D] {d}")
        doc.add_paragraph(f"[CORRECT: {corr}]")
        doc.add_paragraph(f"[HINT: {hint}]")
        doc.add_paragraph("[QUESTION_END]")

    doc.add_paragraph("[EXAM_END]")
    doc.save("test_valid_exam.docx")
    print("✓ Successfully created test_valid_exam.docx (100% valid)")

def create_invalid_exam():
    doc = Document()
    doc.add_paragraph("TÀI LIỆU TEST ĐỀ THI LỖI ĐỊNH DẠNG ĐỂ KIỂM THỬ TÍNH NĂNG ROLLBACK CỦA BACKEND")
    
    # LỖI 1: Thiếu thẻ [EXAM_BEGIN]
    doc.add_paragraph("[EXAM_TITLE: TOEIC READING TEST - MA DE ERROR TEST]")
    doc.add_paragraph("[PART: 5]")

    # Câu 101: Đúng chuẩn
    doc.add_paragraph("[QUESTION: 101]")
    doc.add_paragraph("101. Ms. Carpenter will be attending the conference with ------- marketing team.")
    doc.add_paragraph("[A] she")
    doc.add_paragraph("[B] her")
    doc.add_paragraph("[C] hers")
    doc.add_paragraph("[D] herself")
    doc.add_paragraph("[CORRECT: B]")
    doc.add_paragraph("[HINT: Tính từ sở hữu her.]")
    doc.add_paragraph("[QUESTION_END]")

    # LỖI 2: Sai thứ tự câu hỏi! (Từ câu 101 nhảy cóc lên câu 105)
    doc.add_paragraph("[QUESTION: 105]")
    doc.add_paragraph("105. Câu này bị sai thứ tự (nhảy cóc từ 101 lên 105 mà không có 102, 103, 104).")

    # LỖI 3: Thiếu lựa chọn đáp án [D] (chỉ có A, B, C)
    doc.add_paragraph("[A] Lựa chọn A")
    doc.add_paragraph("[B] Lựa chọn B")
    doc.add_paragraph("[C] Lựa chọn C")

    # LỖI 4: Thiếu thẻ đáp án đúng [CORRECT: ...]
    doc.add_paragraph("[HINT: Câu này cố tình không có thẻ CORRECT và thiếu đáp án D]")

    # LỖI 5: Thiếu thẻ kết thúc câu hỏi [QUESTION_END]

    # LỖI 6: Thiếu thẻ kết thúc toàn bộ bài thi [EXAM_END]
    
    doc.save("test_invalid_exam.docx")
    print("✓ Successfully created test_invalid_exam.docx (Contains 6 critical format errors)")

if __name__ == "__main__":
    create_valid_exam()
    create_invalid_exam()
