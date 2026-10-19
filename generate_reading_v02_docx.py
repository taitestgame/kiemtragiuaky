# -*- coding: utf-8 -*-
"""
Generate complete Reading v02.docx with standard tags, verified answers, hints, and passage images.
"""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from toeic_p5 import PART_5_QUESTIONS
from toeic_p6 import PART_6_GROUPS
from toeic_p7_part1 import PART_7_GROUPS_1
from toeic_p7_part2 import PART_7_GROUPS_2

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = Document()
    
    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("BỘ ĐỀ THI TOEIC READING - READING V02\n(100 CÂU HỎI TIÊU CHUẨN KÈM BỘ THẺ TAGS)")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(26, 54, 93)

    # Tag Specification Table
    p_intro = doc.add_paragraph()
    r = p_intro.add_run("QUY ĐỊNH BỘ THẺ (TAGS) ĐỊNH DẠNG ĐỀ THI TOEIC")
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(43, 108, 176)

    table = doc.add_table(rows=7, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    headers = ["Tên Thẻ (Tag)", "Mục đích sử dụng", "Ví dụ / Giá trị"]
    rows_data = [
        ["[EXAM_BEGIN] / [EXAM_END]", "Đánh dấu bắt đầu và kết thúc toàn bộ đề thi", "[EXAM_TITLE: Reading Test v02]"],
        ["[PART: X]", "Xác định phần thi (5, 6, 7) và tính chất", "[PART: 5] - Incomplete Sentences"],
        ["[SHUFFLE_QUESTIONS: T/F]", "Cho phép / Không cho phép đảo thứ tự câu hỏi", "[SHUFFLE_QUESTIONS: TRUE]"],
        ["[SHUFFLE_ANSWERS: T/F]", "Cho phép / Không cho phép đảo thứ tự các đáp án", "[SHUFFLE_ANSWERS: TRUE]"],
        ["[GROUP_BEGIN] / [GROUP_END]", "Nhóm câu hỏi theo bài đọc (Part 6, Part 7)", "[GROUP_BEGIN: 131-134] ... [GROUP_END]"],
        ["[CORRECT: X] / [HINT: ...]", "Đáp án đúng (A/B/C/D) và gợi ý giải thích chi tiết", "[CORRECT: B] | [HINT: Cần tính từ sở hữu...]"]
    ]

    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = text
        set_cell_background(cell, "2B6CB0")
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, row in enumerate(rows_data):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = text
            if row_idx % 2 == 1:
                set_cell_background(cell, "EDF2F7")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Exam metadata tags
    p_meta = doc.add_paragraph()
    p_meta.add_run("[EXAM_BEGIN]\n").font.bold = True
    p_meta.add_run("[EXAM_TITLE: TOEIC READING TEST - READING V02]\n")
    p_meta.add_run("[EXAM_DURATION: 75]\n")
    p_meta.add_run("[TOTAL_QUESTIONS: 100]\n")
    p_meta.paragraph_format.space_after = Pt(14)

    # ==================== PART 5 ====================
    doc.add_heading("PART 5: INCOMPLETE SENTENCES", level=1)
    p_p5_tag = doc.add_paragraph()
    p_p5_tag.add_run("[PART: 5]\n").font.bold = True
    p_p5_tag.add_run("[PART_TITLE: Incomplete Sentences (Questions 101-130)]\n")
    p_p5_tag.add_run("[SHUFFLE_QUESTIONS: TRUE]\n")
    p_p5_tag.add_run("[SHUFFLE_ANSWERS: TRUE]\n")

    for q in PART_5_QUESTIONS:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(6)
        p_q.paragraph_format.space_after = Pt(2)
        
        r_tag = p_q.add_run(f"[QUESTION: {q['num']}]\n")
        r_tag.font.bold = True
        r_tag.font.color.rgb = RGBColor(43, 108, 176)
        
        r_text = p_q.add_run(f"{q['num']}. {q['text']}\n")
        r_text.font.size = Pt(11)
        
        for opt_key in ["A", "B", "C", "D"]:
            r_opt = p_q.add_run(f"[{opt_key}] ({opt_key}) {q['options'][opt_key]}\n")
            r_opt.font.size = Pt(10.5)
            
        r_ans = p_q.add_run(f"[CORRECT: {q['correct']}]\n")
        r_ans.font.bold = True
        r_ans.font.color.rgb = RGBColor(40, 167, 69)
        
        r_hint = p_q.add_run(f"[HINT: {q['hint']}]\n")
        r_hint.font.italic = True
        r_hint.font.color.rgb = RGBColor(108, 117, 125)
        
        p_q.add_run("[QUESTION_END]\n").font.bold = True

    # ==================== PART 6 ====================
    doc.add_page_break()
    doc.add_heading("PART 6: TEXT COMPLETION", level=1)
    p_p6_tag = doc.add_paragraph()
    p_p6_tag.add_run("[PART: 6]\n").font.bold = True
    p_p6_tag.add_run("[PART_TITLE: Text Completion (Questions 131-146)]\n")
    p_p6_tag.add_run("[SHUFFLE_QUESTIONS: FALSE]\n")
    p_p6_tag.add_run("[SHUFFLE_ANSWERS: TRUE]\n")

    for g in PART_6_GROUPS:
        p_g = doc.add_paragraph()
        p_g.paragraph_format.space_before = Pt(10)
        p_g.add_run(f"[GROUP_BEGIN: {g['group_id']}]\n").font.bold = True
        p_g.add_run(f"[PASSAGE_TITLE: {g['title']}]\n").font.italic = True
        
        p_pass = doc.add_paragraph()
        p_pass.paragraph_format.left_indent = Inches(0.2)
        p_pass.paragraph_format.right_indent = Inches(0.2)
        p_pass.add_run("[PASSAGE_TEXT]\n").font.bold = True
        p_pass.add_run(g['passage_text'] + "\n")
        p_pass.add_run("[/PASSAGE_TEXT]\n").font.bold = True

        for q in g['questions']:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(2)
            
            p_q.add_run(f"[QUESTION: {q['num']}]\n").font.bold = True
            p_q.add_run(f"{q['num']}. {q['text']}\n")
            
            for opt_key in ["A", "B", "C", "D"]:
                p_q.add_run(f"[{opt_key}] ({opt_key}) {q['options'][opt_key]}\n")
                
            r_ans = p_q.add_run(f"[CORRECT: {q['correct']}]\n")
            r_ans.font.bold = True
            r_ans.font.color.rgb = RGBColor(40, 167, 69)
            
            r_hint = p_q.add_run(f"[HINT: {q['hint']}]\n")
            r_hint.font.italic = True
            r_hint.font.color.rgb = RGBColor(108, 117, 125)
            
            p_q.add_run("[QUESTION_END]\n").font.bold = True

        doc.add_paragraph(f"[GROUP_END: {g['group_id']}]\n")

    # ==================== PART 7 ====================
    doc.add_page_break()
    doc.add_heading("PART 7: READING COMPREHENSION", level=1)
    p_p7_tag = doc.add_paragraph()
    p_p7_tag.add_run("[PART: 7]\n").font.bold = True
    p_p7_tag.add_run("[PART_TITLE: Reading Comprehension (Questions 147-200)]\n")
    p_p7_tag.add_run("[SHUFFLE_QUESTIONS: FALSE]\n")
    p_p7_tag.add_run("[SHUFFLE_ANSWERS: TRUE]\n")

    all_p7 = PART_7_GROUPS_1 + PART_7_GROUPS_2
    for g in all_p7:
        p_g = doc.add_paragraph()
        p_g.paragraph_format.space_before = Pt(12)
        p_g.add_run(f"[GROUP_BEGIN: {g['group_id']}]\n").font.bold = True
        p_g.add_run(f"[PASSAGE_TITLE: {g['title']}]\n").font.italic = True
        
        # Add Image if available
        if 'image_file' in g and os.path.exists(g['image_file']):
            p_img_tag = doc.add_paragraph()
            p_img_tag.add_run(f"[PASSAGE_IMAGE: {os.path.basename(g['image_file'])}]\n").font.bold = True
            try:
                doc.add_picture(g['image_file'], width=Inches(5.5))
            except Exception as e:
                print(f"Could not attach image {g['image_file']}: {e}")

        # Add Passage Text
        p_pass = doc.add_paragraph()
        p_pass.paragraph_format.left_indent = Inches(0.2)
        p_pass.paragraph_format.right_indent = Inches(0.2)
        p_pass.add_run("[PASSAGE_TEXT]\n").font.bold = True
        p_pass.add_run(g['passage_text'] + "\n")
        p_pass.add_run("[/PASSAGE_TEXT]\n").font.bold = True

        for q in g['questions']:
            p_q = doc.add_paragraph()
            p_q.paragraph_format.space_before = Pt(4)
            p_q.paragraph_format.space_after = Pt(2)
            
            p_q.add_run(f"[QUESTION: {q['num']}]\n").font.bold = True
            p_q.add_run(f"{q['num']}. {q['text']}\n")
            
            for opt_key in ["A", "B", "C", "D"]:
                p_q.add_run(f"[{opt_key}] ({opt_key}) {q['options'][opt_key]}\n")
                
            r_ans = p_q.add_run(f"[CORRECT: {q['correct']}]\n")
            r_ans.font.bold = True
            r_ans.font.color.rgb = RGBColor(40, 167, 69)
            
            r_hint = p_q.add_run(f"[HINT: {q['hint']}]\n")
            r_hint.font.italic = True
            r_hint.font.color.rgb = RGBColor(108, 117, 125)
            
            p_q.add_run("[QUESTION_END]\n").font.bold = True

        doc.add_paragraph(f"[GROUP_END: {g['group_id']}]\n")

    doc.add_paragraph("\n[EXAM_END]").runs[0].font.bold = True

    output_filename = "Reading v02.docx"
    doc.save(output_filename)
    print(f"Successfully generated {output_filename} with {os.path.getsize(output_filename)} bytes!")

if __name__ == "__main__":
    create_document()
